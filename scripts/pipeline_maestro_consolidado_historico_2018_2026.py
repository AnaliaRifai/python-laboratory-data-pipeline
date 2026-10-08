# -*- coding: utf-8 -*-
"""
PROYECTO: SISTEMA DE CONSULTA HISTÓRICA - LABORATORIO
ETAPA FINAL: CONSOLIDACIÓN HISTÓRICA 2018-2026

Objetivo:
Unificar los consolidados anuales generados con Python en una única base
histórica limpia, validada y lista para análisis.

Entrada:
- consolidado_2018.xlsx
- consolidado_2019.xlsx
- ...
- consolidado_2026.xlsx

Salida:
- consolidado_historico_2018_2026.xlsx
- log_consolidado_historico_2018_2026.xlsx

Criterio general:
Los archivos originales y los consolidados anuales no se modifican.
Este script genera una nueva base final a partir de ellos.
"""

import os
import re
import pandas as pd


# ============================================================
# 1. CONFIGURACIÓN GENERAL
# ============================================================

RAIZ_PROYECTO = r"RUTA_A_TU_PROYECTO\Proyecto Laboratorio"

# IMPORTANTE:
# Para ejecutar este script en otro equipo, reemplazar RAIZ_PROYECTO
# por la ruta local donde se encuentre la carpeta del proyecto.

# Carpeta donde están los consolidados anuales ya validados
CARPETA_CONSOLIDADOS = os.path.join(RAIZ_PROYECTO, "DATOS", "RESPALDOS")

# Carpeta de salida para la base histórica final
CARPETA_SALIDA = os.path.join(RAIZ_PROYECTO, "DATOS", "PROCESADOS")

ANIO_INICIO = 2018
ANIO_FIN = 2026

ARCHIVO_SALIDA = f"consolidado_historico_{ANIO_INICIO}_{ANIO_FIN}.xlsx"
ARCHIVO_LOG = f"log_consolidado_historico_{ANIO_INICIO}_{ANIO_FIN}.xlsx"

os.makedirs(CARPETA_SALIDA, exist_ok=True)


# ============================================================
# 2. CONFIGURACIÓN DE COLUMNAS
# ============================================================

# Columnas que no se incluyen en la base histórica final.
# Se eliminan del dataset final, pero la trazabilidad anual queda documentada
# en el archivo de log.
COLUMNAS_ELIMINAR_FINAL = [
    "archivo",
    "consolidado_origen",
]

# Orden sugerido para la base final.
# Se reemplaza "observaciones" por "cultivo" porque esa columna se usará
# como descripción/cultivo de la muestra.
COLUMNAS_PRIORITARIAS = [
    "anio",
    "muestra",
    "solicitante",
    "localidad",
    "fecha_ingreso",
    "fecha_egreso",
    "cultivo",
    "ms_pct",
    "ph",
    "nt_pct",
    "pb_pct",
    "fdn_pct",
    "fda_pct",
    "lignina_pct",
    "divms_pct",
    "em_mcal_kg_ms",
    "almidon_pct",
    "cnes_pct",
    "extracto_etereo_pct",
    "ns_nt_pct",
    "actividad_ureasica",
    "celulosa_pct",
]

COLUMNAS_TEXTO = [
    "solicitante",
    "localidad",
    "cultivo",
    "actividad_ureasica",
]

COLUMNAS_FECHA = [
    "fecha_ingreso",
    "fecha_egreso",
]

COLUMNAS_NUMERICAS = [
    "anio",
    "muestra",
    "ms_pct",
    "ph",
    "nt_pct",
    "pb_pct",
    "fdn_pct",
    "fda_pct",
    "lignina_pct",
    "divms_pct",
    "em_mcal_kg_ms",
    "almidon_pct",
    "cnes_pct",
    "extracto_etereo_pct",
    "ns_nt_pct",
    "celulosa_pct",
]

# Columnas numéricas que deben quedar con dos decimales.
# No se incluyen anio ni muestra porque son identificadores enteros.
COLUMNAS_DECIMALES = [
    c for c in COLUMNAS_NUMERICAS
    if c not in ("anio", "muestra")
]


# ============================================================
# 3. REGISTROS DE CONTROL
# ============================================================

registro_errores = []
registro_avisos = []
registro_columnas = []
registro_columnas_sin_nombre = []


def registrar_aviso(tipo, detalle):
    registro_avisos.append({
        "tipo": tipo,
        "detalle": str(detalle),
    })
    print(f"Aviso [{tipo}]: {detalle}")


def registrar_error(anio, archivo, error):
    registro_errores.append({
        "anio": anio,
        "archivo": archivo,
        "error": str(error),
    })
    print(f"ERROR {anio} - {archivo}: {error}")


# ============================================================
# 4. FUNCIONES AUXILIARES
# ============================================================

def normalizar_nombre_columna(columna):
    """
    Normaliza encabezados para evitar diferencias menores entre consolidados.
    """
    nombre = str(columna).strip().lower()
    nombre = nombre.replace(" ", "_")
    nombre = re.sub(r"_+", "_", nombre)
    return nombre


def limpiar_texto_columna(serie):
    """
    Limpia columnas de texto:
    - elimina espacios al inicio y al final;
    - convierte cadenas vacías en valores faltantes.
    """
    serie = serie.astype("string")
    serie = serie.str.strip()
    serie = serie.replace("", pd.NA)
    return serie


def estandarizar_cultivo(serie):
    """
    Estandariza la columna cultivo.

    Criterio:
    - se eliminan espacios dobles;
    - se conserva el texto original en cuanto a palabras;
    - se pasa todo a MAYÚSCULAS para evitar duplicados por diferencias
      de escritura como 'silo de maíz', 'Silo de Maíz' o 'SILO DE MAÍZ'.

    Para análisis en Power BI conviene que las categorías de texto
    estén normalizadas.
    """
    serie = limpiar_texto_columna(serie)
    serie = serie.str.replace(r"\s+", " ", regex=True)
    serie = serie.str.upper()
    return serie


def tiene_datos_reales(serie):
    """
    Evalúa si una columna tiene al menos un dato real.
    """
    if serie.empty:
        return False

    valores = serie.astype("string").str.strip()
    return valores.notna().any() and (valores.dropna() != "").any()


def renombrar_observaciones_a_cultivo(df):
    """
    Renombra la columna observaciones a cultivo.

    Si por algún motivo existieran ambas columnas, combina la información
    priorizando cultivo y completando con observaciones cuando cultivo esté vacío.
    """
    if "observaciones" in df.columns and "cultivo" not in df.columns:
        df = df.rename(columns={"observaciones": "cultivo"})

    elif "observaciones" in df.columns and "cultivo" in df.columns:
        cultivo_vacio = df["cultivo"].isna() | (df["cultivo"].astype("string").str.strip() == "")
        df.loc[cultivo_vacio, "cultivo"] = df.loc[cultivo_vacio, "observaciones"]
        df = df.drop(columns=["observaciones"])

    elif "cultivo" not in df.columns:
        df["cultivo"] = pd.NA

    return df



def estandarizar_nombres_finales(df):
    """
    Renombra variantes históricas que quedaron con nombres no estándar
    en los consolidados anuales.

    Ejemplo:
    'celulosa' -> 'celulosa_pct'
    """
    renombres = {
        "celulosa": "celulosa_pct",
    }

    for origen, destino in renombres.items():
        if origen in df.columns and destino not in df.columns:
            df = df.rename(columns={origen: destino})

        elif origen in df.columns and destino in df.columns:
            # Si existen ambas, se completa destino con origen cuando destino esté vacío.
            destino_vacio = df[destino].isna()
            df.loc[destino_vacio, destino] = df.loc[destino_vacio, origen]
            df = df.drop(columns=[origen])

    return df


def integrar_columnas_sin_nombre(df, anio, archivo_origen):
    """
    Trata columnas generadas por encabezados vacíos: sin_nombre_...

    Criterio:
    - si están totalmente vacías, se eliminan;
    - si tienen contenido, se anexan a cultivo para no perder información;
    - en ambos casos queda registro en el log.
    """
    columnas_sin_nombre = [
        c for c in df.columns
        if str(c).startswith("sin_nombre")
    ]

    if not columnas_sin_nombre:
        return df

    if "cultivo" not in df.columns:
        df["cultivo"] = pd.NA

    for col in columnas_sin_nombre:
        con_datos = tiene_datos_reales(df[col])

        if not con_datos:
            registro_columnas_sin_nombre.append({
                "anio": anio,
                "archivo_consolidado": archivo_origen,
                "columna": col,
                "accion": "eliminada por estar vacía",
                "filas_con_dato": 0,
            })
            df = df.drop(columns=[col])
            continue

        mascara = df[col].notna() & (df[col].astype("string").str.strip() != "")
        filas_con_dato = int(mascara.sum())

        cultivo = df.loc[mascara, "cultivo"].astype("string").fillna("").str.strip()
        extra = df.loc[mascara, col].astype("string").str.strip()

        df.loc[mascara, "cultivo"] = (
            cultivo.where(cultivo == "", cultivo + " | ") + extra
        )

        registro_columnas_sin_nombre.append({
            "anio": anio,
            "archivo_consolidado": archivo_origen,
            "columna": col,
            "accion": "contenido anexado a cultivo y columna eliminada",
            "filas_con_dato": filas_con_dato,
        })

        df = df.drop(columns=[col])

    return df


def estandarizar_tipos(df):
    """
    Estandariza tipos de datos para el consolidado histórico.
    """
    # Fechas
    for col in COLUMNAS_FECHA:
        if col in df.columns:
            df[col] = pd.to_datetime(df[col], errors="coerce", dayfirst=False)

    # Numéricas
    for col in COLUMNAS_NUMERICAS:
        if col in df.columns:
            if col in ("anio", "muestra"):
                df[col] = pd.to_numeric(df[col], errors="coerce").astype("Int64")
            else:
                df[col] = pd.to_numeric(df[col], errors="coerce")

    # Redondear variables analíticas a dos decimales
    for col in COLUMNAS_DECIMALES:
        if col in df.columns:
            df[col] = df[col].round(2)

    # Texto
    for col in COLUMNAS_TEXTO:
        if col in df.columns and col != "cultivo":
            df[col] = limpiar_texto_columna(df[col])

    if "cultivo" in df.columns:
        df["cultivo"] = estandarizar_cultivo(df["cultivo"])

    return df


def ordenar_columnas(df):
    """
    Ordena las columnas para que la base final tenga una estructura estable.
    """
    columnas_inicio = [
        c for c in COLUMNAS_PRIORITARIAS
        if c in df.columns
    ]

    otras_columnas = [
        c for c in df.columns
        if c not in columnas_inicio
        and c not in COLUMNAS_ELIMINAR_FINAL
    ]

    return df[columnas_inicio + otras_columnas]


def aplicar_formato_excel(ruta_excel):
    """
    Aplica formato visual al Excel final:
    - fechas como DD/MM/YYYY;
    - variables analíticas con dos decimales;
    - filtros;
    - panel congelado;
    - anchos de columna aproximados.
    """
    from openpyxl import load_workbook

    wb = load_workbook(ruta_excel)

    if "base_historica" in wb.sheetnames:
        ws = wb["base_historica"]

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        encabezados = {
            cell.value: cell.column
            for cell in ws[1]
        }

        # Formato fechas
        for col_fecha in COLUMNAS_FECHA:
            if col_fecha in encabezados:
                col_idx = encabezados[col_fecha]
                for fila in range(2, ws.max_row + 1):
                    ws.cell(row=fila, column=col_idx).number_format = "DD/MM/YYYY"

        # Formato números con dos decimales
        for col_num in COLUMNAS_DECIMALES:
            if col_num in encabezados:
                col_idx = encabezados[col_num]
                for fila in range(2, ws.max_row + 1):
                    ws.cell(row=fila, column=col_idx).number_format = "0.00"

        # Formato enteros
        for col_entera in ("anio", "muestra"):
            if col_entera in encabezados:
                col_idx = encabezados[col_entera]
                for fila in range(2, ws.max_row + 1):
                    ws.cell(row=fila, column=col_idx).number_format = "0"

        # Ancho de columnas
        for col in ws.columns:
            letra = col[0].column_letter
            encabezado = col[0].value

            if encabezado in ("solicitante", "localidad", "cultivo"):
                ws.column_dimensions[letra].width = 30
            elif encabezado in COLUMNAS_FECHA:
                ws.column_dimensions[letra].width = 14
            else:
                ws.column_dimensions[letra].width = 13

    wb.save(ruta_excel)


# ============================================================
# 5. LECTURA Y UNIÓN DE CONSOLIDADOS ANUALES
# ============================================================

tablas = []

print("\n--- LECTURA DE CONSOLIDADOS ANUALES ---")

for anio in range(ANIO_INICIO, ANIO_FIN + 1):

    nombre_archivo = f"consolidado_{anio}.xlsx"
    ruta_archivo = os.path.join(CARPETA_CONSOLIDADOS, nombre_archivo)

    if not os.path.exists(ruta_archivo):
        registrar_error(
            anio,
            nombre_archivo,
            "No se encontró el consolidado anual en la carpeta esperada"
        )
        continue

    try:
        df = pd.read_excel(ruta_archivo)

        # Normalizar encabezados
        df.columns = [normalizar_nombre_columna(c) for c in df.columns]

        # Registrar estructura original anual en el log
        registro_columnas.append({
            "anio": anio,
            "archivo_consolidado": nombre_archivo,
            "cantidad_filas": len(df),
            "cantidad_columnas": len(df.columns),
            "columnas": ", ".join(df.columns.tolist()),
        })

        # Agregar año de origen
        df["anio"] = anio

        # Renombrar observaciones a cultivo antes de controles adicionales
        df = renombrar_observaciones_a_cultivo(df)

        # Estandarizar nombres históricos que hayan quedado con variantes
        df = estandarizar_nombres_finales(df)

        # Tratar columnas sin nombre
        df = integrar_columnas_sin_nombre(df, anio, nombre_archivo)

        # Estandarizar tipos dentro de cada año
        df = estandarizar_tipos(df)

        tablas.append(df)

        print(f"OK {anio}: {nombre_archivo} | Filas: {len(df)} | Columnas: {len(df.columns)}")

    except Exception as e:
        registrar_error(anio, nombre_archivo, e)


if not tablas:
    raise SystemExit("No se pudo leer ningún consolidado anual. Revisar rutas y errores.")


# ============================================================
# 6. CONSOLIDADO HISTÓRICO
# ============================================================

consolidado = pd.concat(tablas, ignore_index=True, sort=False)

# Eliminar columnas que no se incluyen en la base histórica final
columnas_a_eliminar = [
    col for col in COLUMNAS_ELIMINAR_FINAL
    if col in consolidado.columns
]

if columnas_a_eliminar:
    registrar_aviso(
        "columnas eliminadas de la base final",
        ", ".join(columnas_a_eliminar)
    )
    consolidado = consolidado.drop(columns=columnas_a_eliminar)

# Eliminar columnas totalmente vacías luego de unir todos los años
columnas_vacias = [
    col for col in consolidado.columns
    if consolidado[col].isna().all()
]

if columnas_vacias:
    registrar_aviso(
        "columnas vacías eliminadas",
        ", ".join(columnas_vacias)
    )
    consolidado = consolidado.drop(columns=columnas_vacias)

# Estandarizar nombres finales por seguridad luego de unir todo
consolidado = estandarizar_nombres_finales(consolidado)

# Reaplicar tipos luego de unir todo
consolidado = estandarizar_tipos(consolidado)

# Ordenar columnas finales
consolidado = ordenar_columnas(consolidado)


# ============================================================
# 7. CONTROLES DE CALIDAD HISTÓRICOS
# ============================================================

print("\n--- CONTROLES DE CALIDAD HISTÓRICOS ---")

resumen_por_anio = (
    consolidado
    .groupby("anio", dropna=False)
    .agg(
        muestras=("muestra", "count"),
        fecha_ingreso_min=("fecha_ingreso", "min"),
        fecha_ingreso_max=("fecha_ingreso", "max"),
    )
    .reset_index()
)

print("\nResumen por año:")
print(resumen_por_anio)

duplicados = consolidado[
    consolidado.duplicated(subset=["muestra"], keep=False)
].sort_values(["muestra", "anio"])

print("\nDuplicados históricos por muestra:", len(duplicados))

fechas_invertidas = consolidado[
    consolidado["fecha_egreso"].notna()
    & consolidado["fecha_ingreso"].notna()
    & (consolidado["fecha_egreso"] < consolidado["fecha_ingreso"])
].copy()

print("Fechas de egreso anteriores al ingreso:", len(fechas_invertidas))

# Se permite 2017 porque algunos informes de 2018 tienen fecha de ingreso de fines de 2017.
fecha_minima_esperada = pd.Timestamp("2017-01-01")
fecha_maxima_esperada = pd.Timestamp(f"{ANIO_FIN}-12-31")

fechas_fuera_periodo = consolidado[
    consolidado["fecha_ingreso"].notna()
    & (
        (consolidado["fecha_ingreso"] < fecha_minima_esperada)
        | (consolidado["fecha_ingreso"] > fecha_maxima_esperada)
    )
].copy()

print("Fechas de ingreso fuera del período esperado:", len(fechas_fuera_periodo))

columnas_sin_nombre_df = pd.DataFrame(registro_columnas_sin_nombre)

print("Columnas sin nombre tratadas:", len(columnas_sin_nombre_df))

nulos_por_columna = (
    consolidado
    .isna()
    .sum()
    .rename("valores_nulos")
    .to_frame()
    .reset_index()
    .rename(columns={"index": "columna"})
)

nulos_por_columna["total_filas"] = len(consolidado)
nulos_por_columna["porcentaje_nulos"] = (
    nulos_por_columna["valores_nulos"] / len(consolidado) * 100
).round(2)

columnas_analiticas = [
    c for c in COLUMNAS_DECIMALES
    if c in consolidado.columns
]

cobertura_analiticas = []

for col in columnas_analiticas:
    cobertura_analiticas.append({
        "variable": col,
        "valores_con_dato": int(consolidado[col].notna().sum()),
        "valores_faltantes": int(consolidado[col].isna().sum()),
        "porcentaje_con_dato": round(consolidado[col].notna().mean() * 100, 2),
    })

cobertura_analiticas = pd.DataFrame(cobertura_analiticas)


# ============================================================
# 8. EXPORTACIÓN
# ============================================================

ruta_salida = os.path.join(CARPETA_SALIDA, ARCHIVO_SALIDA)
ruta_log = os.path.join(CARPETA_SALIDA, ARCHIVO_LOG)

print("\n--- EXPORTACIÓN ---")

with pd.ExcelWriter(
    ruta_salida,
    engine="openpyxl",
    date_format="DD/MM/YYYY",
    datetime_format="DD/MM/YYYY"
) as writer:
    consolidado.to_excel(
        writer,
        sheet_name="base_historica",
        index=False
    )

aplicar_formato_excel(ruta_salida)

with pd.ExcelWriter(
    ruta_log,
    engine="openpyxl",
    date_format="DD/MM/YYYY",
    datetime_format="DD/MM/YYYY"
) as writer:

    resumen_por_anio.to_excel(
        writer,
        sheet_name="resumen_por_anio",
        index=False
    )

    pd.DataFrame(registro_columnas).to_excel(
        writer,
        sheet_name="columnas_por_anio",
        index=False
    )

    nulos_por_columna.to_excel(
        writer,
        sheet_name="nulos_por_columna",
        index=False
    )

    cobertura_analiticas.to_excel(
        writer,
        sheet_name="cobertura_analiticas",
        index=False
    )

    duplicados.to_excel(
        writer,
        sheet_name="duplicados_muestra",
        index=False
    )

    fechas_invertidas.to_excel(
        writer,
        sheet_name="fechas_invertidas",
        index=False
    )

    fechas_fuera_periodo.to_excel(
        writer,
        sheet_name="fechas_fuera_periodo",
        index=False
    )

    columnas_sin_nombre_df.to_excel(
        writer,
        sheet_name="columnas_sin_nombre",
        index=False
    )

    pd.DataFrame(registro_avisos).to_excel(
        writer,
        sheet_name="avisos",
        index=False
    )

    pd.DataFrame(registro_errores).to_excel(
        writer,
        sheet_name="errores",
        index=False
    )


# ============================================================
# 9. RESUMEN FINAL
# ============================================================

print("\n--- RESUMEN FINAL ---")
print("Archivos anuales esperados:", ANIO_FIN - ANIO_INICIO + 1)
print("Archivos anuales leídos:", len(tablas))
print("Errores de lectura:", len(registro_errores))
print("Filas consolidadas:", len(consolidado))
print("Columnas finales:", len(consolidado.columns))
print("Duplicados históricos por muestra:", len(duplicados))
print("Fechas invertidas:", len(fechas_invertidas))
print("Fechas fuera del período esperado:", len(fechas_fuera_periodo))
print("Columnas sin nombre tratadas:", len(columnas_sin_nombre_df))

print("\nArchivo final para análisis:")
print(ruta_salida)

print("\nLog de control histórico:")
print(ruta_log)

print("\nProceso finalizado.")
